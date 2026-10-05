"""Controlled spreadsheet migration: stage -> validate -> de-duplicate -> (dry-run | migrate) -> reconcile -> rollback.

Rules that protect existing data (NFR 01):
  * existing records are NEVER updated or deleted by an import;
  * a database backup is taken before any committed import;
  * every created row carries import_batch so the batch can be rolled back precisely.
"""
import csv
import re
from datetime import datetime
from pathlib import Path

from django.conf import settings
from django.db import transaction
from django.utils import timezone

from catalog import search
from catalog.models import Biblio, Item
from core import audit
from core.ops import backup
from migration_tool.models import ImportBatch, StagingRow

COLUMNS = ["accession_no", "title", "author", "isbn", "call_number", "publisher", "year", "shelf_location", "is_reference"]
CHUNK = 2000


class ReconciliationError(Exception):
    pass


class SimulatedFailure(Exception):
    pass


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            for i, row in enumerate(csv.DictReader(fh), start=2):
                yield i, {k.strip().lower(): (v or "").strip() for k, v in row.items() if k}
        return
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    for i, values in enumerate(rows, start=2):
        if all(v in (None, "") for v in values):
            continue
        yield i, {header[j]: ("" if v is None else str(v).strip()) for j, v in enumerate(values) if j < len(header) and header[j]}
    wb.close()


def validate(d):
    errors = []
    if not d.get("accession_no"):
        errors.append("accession_no is required")
    elif not re.fullmatch(r"[A-Za-z0-9._-]{1,40}", d["accession_no"]):
        errors.append("accession_no has invalid characters")
    if not d.get("title"):
        errors.append("title is required")
    elif len(d["title"]) > 300:
        errors.append("title longer than 300 characters")
    isbn = re.sub(r"[-\s]", "", d.get("isbn", ""))
    if isbn and not re.fullmatch(r"\d{9}[\dXx]|\d{13}", isbn):
        errors.append("isbn is not 10 or 13 digits")
    year = d.get("year", "")
    if year:
        y = year.split(".")[0]
        if not y.isdigit() or not (1000 <= int(y) <= datetime.now().year + 1):
            errors.append("year is not a valid year")
    if d.get("is_reference", "").lower() not in ("", "y", "n", "yes", "no", "true", "false", "1", "0"):
        errors.append("is_reference must be yes/no")
    return errors


def _write_reports(batch):
    folder = settings.DATA_DIR / "reports"
    folder.mkdir(parents=True, exist_ok=True)
    err_path = folder / f"batch_{batch.pk}_errors.csv"
    with open(err_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["row_no", "status", "accession_no", "title", "problem"])
        for r in batch.rows.exclude(status__in=["VALID", "MIGRATED"]).order_by("row_no").iterator():
            w.writerow([r.row_no, r.status, r.data.get("accession_no", ""), r.data.get("title", ""), r.error])
    rec = reconcile(batch)
    rec_path = folder / f"batch_{batch.pk}_reconciliation.txt"
    rec_path.write_text("\n".join(f"{k}: {v}" for k, v in rec.items()) + "\n", encoding="utf-8")
    return err_path, rec_path


def reconcile(batch):
    target = Item.objects.filter(import_batch=batch).count()
    accounted = batch.migrated_rows + batch.rejected_rows + batch.duplicate_rows
    if batch.dry_run:
        accounted = batch.valid_rows + batch.rejected_rows + batch.duplicate_rows
    return {
        "batch": batch.pk, "status": batch.status, "source_rows": batch.total_rows, "valid": batch.valid_rows,
        "rejected": batch.rejected_rows, "duplicates": batch.duplicate_rows, "migrated": batch.migrated_rows,
        "target_items_in_batch": target,
        "source_equals_valid_plus_rejected_plus_duplicates": batch.total_rows == accounted,
        "target_matches_migrated": target == batch.migrated_rows,
    }


def _stage(batch, path):
    existing = set(Item.objects.values_list("accession_no", flat=True)) | set(Item.objects.values_list("barcode", flat=True))
    seen, buffer = set(), []
    counts = {"total": 0, "valid": 0, "rejected": 0, "duplicate": 0}
    for row_no, data in read_rows(path):
        counts["total"] += 1
        errors = validate(data)
        status, msg = "VALID", ""
        if errors:
            status, msg = "REJECTED", "; ".join(errors)
        elif data["accession_no"] in existing:
            status, msg = "DUPLICATE", "accession_no already exists in the library database (existing record kept unchanged)"
        elif data["accession_no"] in seen:
            status, msg = "DUPLICATE", "accession_no repeated earlier in the file (first occurrence kept)"
        if status == "VALID":
            seen.add(data["accession_no"])
        counts[{"VALID": "valid", "REJECTED": "rejected", "DUPLICATE": "duplicate"}[status]] += 1
        buffer.append(StagingRow(batch=batch, row_no=row_no, data=data, status=status, error=msg))
        if len(buffer) >= CHUNK:
            StagingRow.objects.bulk_create(buffer)
            buffer = []
    if buffer:
        StagingRow.objects.bulk_create(buffer)
    batch.total_rows, batch.valid_rows = counts["total"], counts["valid"]
    batch.rejected_rows, batch.duplicate_rows = counts["rejected"], counts["duplicate"]
    batch.save()


def _load_chunk(batch, rows):
    biblios = []
    for r in rows:
        d = r.data
        year = d.get("year", "").split(".")[0]
        biblios.append(Biblio(title=d["title"], author=d.get("author", "")[:200], isbn=re.sub(r"[-\s]", "", d.get("isbn", ""))[:20],
                              call_number=d.get("call_number", "")[:60], publisher=d.get("publisher", "")[:200],
                              year=int(year) if year else None, source_key=d["accession_no"], import_batch=batch))
    Biblio.objects.bulk_create(biblios)
    keys = [r.data["accession_no"] for r in rows]
    mapping = dict(Biblio.objects.filter(import_batch=batch, source_key__in=keys).values_list("source_key", "id"))
    items = [Item(biblio_id=mapping[r.data["accession_no"]], accession_no=r.data["accession_no"], barcode=r.data["accession_no"],
                  shelf_location=r.data.get("shelf_location", "")[:40], import_batch=batch,
                  is_reference=r.data.get("is_reference", "").lower() in ("y", "yes", "true", "1")) for r in rows]
    Item.objects.bulk_create(items)
    search.reindex_many(Biblio.objects.filter(import_batch=batch, source_key__in=keys))
    StagingRow.objects.filter(pk__in=[r.pk for r in rows]).update(status="MIGRATED")
    return len(items)


def run_import(path, dry_run=True, actor="cli", simulate_failure=False):
    path = Path(path)
    batch = ImportBatch.objects.create(source_file=path.name, dry_run=dry_run, status=ImportBatch.Status.FAILED)
    audit.log(actor, "IMPORT_STARTED", batch.pk, f"{path.name} dry_run={dry_run}")
    try:
        _stage(batch, path)
        if dry_run:
            batch.status = ImportBatch.Status.DRY_RUN
        else:
            batch.backup_file = backup.create_backup(f"pre-import-{batch.pk}").name  # NFR 01: back up first
            batch.save()
            valid = list(batch.rows.filter(status="VALID").order_by("row_no"))
            chunks = [valid[i:i + CHUNK] for i in range(0, len(valid), CHUNK)]
            migrated = 0
            if simulate_failure:  # chunked, NON-atomic: leaves partial data so the restore drill is meaningful
                for n, chunk in enumerate(chunks):
                    with transaction.atomic():
                        migrated += _load_chunk(batch, chunk)
                    if n >= 1 or n == len(chunks) - 1:
                        raise SimulatedFailure("simulated crash after partial load")
            else:
                with transaction.atomic():
                    for chunk in chunks:
                        migrated += _load_chunk(batch, chunk)
                    batch.migrated_rows = migrated
                    rec = reconcile(batch)
                    if not (rec["source_equals_valid_plus_rejected_plus_duplicates"] and rec["target_matches_migrated"]):
                        raise ReconciliationError(f"source/target counts do not reconcile: {rec}")
            batch.migrated_rows = migrated
            batch.status = ImportBatch.Status.COMMITTED
        batch.finished = timezone.now()
        batch.save()
        _write_reports(batch)
        audit.log(actor, "IMPORT_FINISHED", batch.pk, f"status={batch.status} migrated={batch.migrated_rows}")
        return batch
    except Exception as exc:
        batch.status = ImportBatch.Status.FAILED
        batch.error = f"{exc.__class__.__name__}: {exc}"[:300]
        batch.finished = timezone.now()
        batch.save()
        audit.log(actor, "IMPORT_FAILED", batch.pk, batch.error)
        raise


@transaction.atomic
def rollback_batch(batch_id, actor="cli"):
    batch = ImportBatch.objects.get(pk=batch_id)
    if batch.status == ImportBatch.Status.ROLLED_BACK:
        return batch
    Item.objects.filter(import_batch=batch).delete()
    Biblio.objects.filter(import_batch=batch).delete()
    search.purge_orphans()
    batch.status, batch.migrated_rows = ImportBatch.Status.ROLLED_BACK, 0
    batch.save()
    audit.log(actor, "IMPORT_ROLLED_BACK", batch.pk, "all rows from this batch removed")
    return batch
